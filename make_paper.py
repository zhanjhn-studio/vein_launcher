"""生成 Vein Launcher 专用纸（A4 空白页 + 左上角 Logo 页眉）"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

LOGO = Path("assets/logo.png")
OUT = Path("VeinLauncher_专用纸.docx")


def add_bottom_border(paragraph, color="E8ECF3", size="6"):
    p_pr = paragraph._element.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.8)
    section.bottom_margin = Cm(2.0)
    section.header_distance = Cm(1.1)
    section.footer_distance = Cm(1.0)

    # 页眉：左上角 Logo + 右侧标识，底部细线
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.add_run().add_picture(str(LOGO), height=Cm(0.95))
    tag = header.add_run("　　VEIN LAUNCHER")
    tag.font.size = Pt(8.5)
    tag.font.color.rgb = RGBColor(0x8C, 0x97, 0xA8)
    tag.font.name = "微软雅黑"
    tag._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    add_bottom_border(header)

    # 页脚：居中页码
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    dash = footer.add_run("— ")
    dash.font.size = Pt(8)
    dash.font.color.rgb = RGBColor(0xB4, 0xBC, 0xC8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._element.append(fld)
    tail = footer.add_run(" —")
    tail.font.size = Pt(8)
    tail.font.color.rgb = RGBColor(0xB4, 0xBC, 0xC8)

    # 正文：整页留白
    body = doc.add_paragraph()
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(0)

    doc.save(str(OUT))
    print("paper:", OUT, OUT.stat().st_size)


if __name__ == "__main__":
    build()
